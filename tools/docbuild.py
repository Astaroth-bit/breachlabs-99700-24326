"""Shared helpers for building ePortfolio docx files in the style of Criterion B.i."""
import copy, docx
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_BREAK

TEMPLATE = '/tmp/claude-0/-home-user-breachlabs-99700-24326/48f6b551-31b0-51b5-a4d5-483d74f26dad/scratchpad/Criterion_B_i_Design_Specifications_v3.docx'
TEAL = "1F4E5F"; BORDER = "8FA3AA"
FULL = 10512  # dxa

def new_doc():
    d = docx.Document(TEMPLATE)
    body = d.element.body
    for el in list(body):
        if el.tag != qn('w:sectPr'):
            body.remove(el)
    return d

def _run(p, text, size=10.5, bold=False, italic=False, color=None):
    r = p.add_run(text)
    r.font.name = "Arial"; r._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial')
    r.font.size = Pt(size); r.bold = bold or None; r.italic = italic or None
    if color: r.font.color.rgb = RGBColor.from_string(color)
    return r

def _sty(d, name):
    for st in d.styles:
        if st.name and st.name.lower() == name.lower(): return st
    return None

def H1(d, t):
    p = d.add_paragraph(style=_sty(d, "Heading 1")); _run(p, t, 16, True, color=TEAL); return p
def H2(d, t):
    p = d.add_paragraph(style=_sty(d, "Heading 2")); _run(p, t, 12, True, color=TEAL); return p
def H3(d, t):
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    _run(p, t, 11, True, color=TEAL); return p

def P(d, text, size=10.5, after=6, italic=False):
    """text may contain **bold** segments."""
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    parts = text.split("**")
    for i, s in enumerate(parts):
        if s: _run(p, s, size, bold=(i % 2 == 1), italic=italic)
    return p

def bullets(d, items, size=10.5):
    for it in items:
        p = d.add_paragraph(); p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.18); p.paragraph_format.space_after = Pt(3)
        parts = ("•  " + it).split("**")
        for i, s in enumerate(parts):
            if s: _run(p, s, size, bold=(i % 2 == 1))

def page_break(d):
    p = d.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)

def image(d, path, width_in=7.3, caption=None):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = bool(caption)
    p.add_run().add_picture(path, width=Inches(width_in))
    if caption:
        c = d.add_paragraph(); c.paragraph_format.space_after = Pt(8)
        _run(c, caption, 9, italic=True, color="33444B")

def _set_cell(cell, width, fill=None):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ('w:tcW', 'w:tcBorders', 'w:shd', 'w:tcMar'):
        for e in tcPr.findall(qn(tag)): tcPr.remove(e)
    w = OxmlElement('w:tcW'); w.set(qn('w:type'), 'dxa'); w.set(qn('w:w'), str(width)); tcPr.append(w)
    b = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:val'), 'single'); e.set(qn('w:color'), BORDER); e.set(qn('w:sz'), '4'); b.append(e)
    tcPr.append(b)
    if fill:
        s = OxmlElement('w:shd'); s.set(qn('w:fill'), fill); s.set(qn('w:color'), 'auto'); s.set(qn('w:val'), 'clear'); tcPr.append(s)
    m = OxmlElement('w:tcMar')
    for side, v in (('top', 70), ('left', 90), ('bottom', 70), ('right', 90)):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:type'), 'dxa'); e.set(qn('w:w'), str(v)); m.append(e)
    tcPr.append(m)

def _fill_cell(cell, content, label=False, size=8.5, bold=False, color=None):
    cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p) if cell.paragraphs else None
    paras = content if isinstance(content, list) else str(content).split("\n")
    for i, para in enumerate(paras):
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(4 if i < len(paras) - 1 else 0)
        parts = para.split("**")
        for j, s in enumerate(parts):
            if not s: continue
            if label: _run(p, s, 9, True, color="FFFFFF")
            else: _run(p, s, size, bold=bold or (j % 2 == 1), color=color)

def table(d, rows, widths, header=None, label_col=True, size=8.5, fills=None, repeat_header=True):
    """rows: list of lists of str. fills: dict (r,c)->hex for body cells."""
    ncol = len(widths)
    t = d.add_table(rows=0, cols=ncol)
    tblPr = t._tbl.tblPr
    w = OxmlElement('w:tblW'); w.set(qn('w:type'), 'dxa'); w.set(qn('w:w'), str(sum(widths))); tblPr.append(w)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
    grid = t._tbl.tblGrid
    for gc, wv in zip(grid.findall(qn('w:gridCol')), widths): gc.set(qn('w:w'), str(wv))
    def add_row(vals, is_header=False, ri=None):
        r = t.add_row()
        trPr = r._tr.get_or_add_trPr()
        cs = OxmlElement('w:cantSplit'); trPr.append(cs)
        if is_header and repeat_header:
            th = OxmlElement('w:tblHeader'); trPr.append(th)
        for ci, (cell, val) in enumerate(zip(r.cells, vals)):
            lab = is_header or (label_col and ci == 0)
            fill = TEAL if lab else (fills.get((ri, ci)) if fills else None)
            _set_cell(cell, widths[ci], fill)
            _fill_cell(cell, val, label=lab, size=size)
    if header: add_row(header, True)
    for ri, row in enumerate(rows): add_row(row, False, ri)
    after = d.add_paragraph(); after.paragraph_format.space_after = Pt(4)
    return t

def code(d, text, size=8):
    t = table(d, [[text]], [FULL], label_col=False, size=size)
    for p in t.rows[0].cells[0].paragraphs:
        for r in p.runs:
            r.font.name = "Courier New"; r._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Courier New')
    _set_cell(t.rows[0].cells[0], FULL, "F7F3EA")
    return t
